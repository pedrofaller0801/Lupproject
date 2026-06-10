"""
Pipeline de ingestão de documentos para a base de contexto.

Uso via CLI:
    python backend/ingestao.py --tipo manual
    python backend/ingestao.py --tipo referencia
"""

import argparse
import sys
import time
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from google import genai
from google.genai import types as genai_types
import tiktoken
from supabase import create_client

from config import (
    DIR_MANUAL,
    DIR_PROJETOS,
    GEMINI_API_KEY,
    MIN_CHARS_PAGINA,
    OVERLAP_CHUNK,
    SUPABASE_KEY,
    SUPABASE_URL,
    TAMANHO_CHUNK,
    validar_config,
)
from visao import pdf_para_imagens

# Encoding para contagem de tokens (boa aproximação para português)
_ENCODING = tiktoken.get_encoding("cl100k_base")

# Cliente Gemini — inicializado por configurar_gemini()
_GEMINI_CLIENT: genai.Client | None = None


# ---------------------------------------------------------------------------
# Funções auxiliares
# ---------------------------------------------------------------------------

def configurar_gemini() -> None:
    """Inicializa o cliente Gemini com a chave de API (endpoint v1 estável)."""
    global _GEMINI_CLIENT
    _GEMINI_CLIENT = genai.Client(
        api_key=GEMINI_API_KEY,
        http_options={"api_version": "v1"},
    )


def _get_client() -> genai.Client:
    """Retorna o cliente Gemini, inicializando-o se necessário."""
    if _GEMINI_CLIENT is None:
        configurar_gemini()
    return _GEMINI_CLIENT


def gerar_embedding_documento(texto: str, max_tentativas: int = 5) -> list[float]:
    """
    Gera embedding para armazenamento (task_type=RETRIEVAL_DOCUMENT).
    Faz retry automático em caso de 429, aguardando o retryDelay sugerido pela API.
    """
    import re as _re

    for tentativa in range(max_tentativas):
        try:
            resultado = _get_client().models.embed_content(
                model="gemini-embedding-001",
                contents=texto,
                config=genai_types.EmbedContentConfig(
                    task_type="RETRIEVAL_DOCUMENT",
                    output_dimensionality=768,
                ),
            )
            time.sleep(1.0)  # pausa mínima entre chamadas
            return resultado.embeddings[0].values

        except Exception as e:
            msg = str(e)
            eh_transitorio = any(c in msg for c in ("429", "RESOURCE_EXHAUSTED"))
            ultima = tentativa == max_tentativas - 1

            if not eh_transitorio or ultima:
                raise

            match = _re.search(r"retryDelay.*?(\d+)s", msg)
            espera = int(match.group(1)) + 3 if match else 60 * (tentativa + 1)
            print(f"  [429] Embedding throttled — aguardando {espera}s (tentativa {tentativa + 1}/{max_tentativas})...")
            time.sleep(espera)


def extrair_texto_paginas(caminho_pdf: Path) -> list[tuple[int, str]]:
    """
    Extrai o texto de cada página de um PDF.

    Retorna lista de (numero_pagina, texto).
    Páginas com menos de MIN_CHARS_PAGINA caracteres são marcadas como gráficas
    e receberão aviso no terminal.
    """
    paginas: list[tuple[int, str]] = []

    with fitz.open(str(caminho_pdf)) as doc:
        for i, pagina in enumerate(doc, start=1):
            texto = pagina.get_text("text").strip()
            if len(texto) < MIN_CHARS_PAGINA:
                print(
                    f"  [AVISO] Página {i}: {len(texto)} chars — "
                    "considerada gráfica, ignorada no chunking de texto."
                )
            paginas.append((i, texto))

    return paginas


def extrair_texto_docx(caminho: Path) -> str:
    """
    Extrai texto de um arquivo Word (.docx), incluindo parágrafos e tabelas.
    """
    doc = Document(str(caminho))
    partes: list[str] = []

    for paragrafo in doc.paragraphs:
        txt = paragrafo.text.strip()
        if txt:
            partes.append(txt)

    for tabela in doc.tables:
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                partes.append(" | ".join(celulas))

    return "\n\n".join(partes)


def chunk_texto(texto: str) -> list[str]:
    """
    Divide um texto em chunks de TAMANHO_CHUNK tokens com OVERLAP_CHUNK de sobreposição.
    Retorna lista de strings.
    """
    tokens = _ENCODING.encode(texto)
    if not tokens:
        return []

    chunks: list[str] = []
    passo = TAMANHO_CHUNK - OVERLAP_CHUNK

    for inicio in range(0, len(tokens), passo):
        fatia = tokens[inicio : inicio + TAMANHO_CHUNK]
        chunk = _ENCODING.decode(fatia).strip()
        if chunk:
            chunks.append(chunk)

    return chunks


# ---------------------------------------------------------------------------
# Pipeline principal
# ---------------------------------------------------------------------------

def indexar_documento(caminho: Path, tipo: str, grupo: str | None = None) -> int:
    """
    Indexa um único PDF ou DOCX na base de contexto do Supabase.

    Fluxo:
        1. Extrai texto (PyMuPDF para PDF, python-docx para DOCX)
        2. Para referências PDF, converte páginas em imagens JPEG (cache local)
        3. Quebra em chunks de 300 tokens / overlap 50
        4. Gera embeddings com gemini-embedding-001
        5. Salva no Supabase com tipo, grupo e metadados

    Args:
        caminho: Caminho para o arquivo PDF ou DOCX.
        tipo:    'manual' ou 'referencia'.
        grupo:   Nome do grupo organizacional (opcional, ex: 'Churrasqueiras').

    Returns:
        Número de chunks salvos no banco.
    """
    supabase = create_client(SUPABASE_URL, SUPABASE_KEY)
    ext = caminho.suffix.lower()

    print(f"\n{'─'*60}")
    print(f"Indexando: {caminho.name}  (tipo: {tipo})")
    print(f"{'─'*60}")

    # Etapa 1 — Extração de texto (branch por tipo de arquivo)
    if ext == ".docx":
        texto_total = extrair_texto_docx(caminho)
        if len(texto_total) < MIN_CHARS_PAGINA:
            print("  [AVISO] Nenhum texto extraível no DOCX. Pulando.")
            return 0
    else:  # .pdf
        paginas = extrair_texto_paginas(caminho)

        # Etapa 2 — Para referências PDF, gera cache de imagens
        if tipo == "referencia":
            print("  Convertendo páginas em imagens (300 DPI)...")
            imagens_geradas = pdf_para_imagens(caminho)
            print(f"  {len(imagens_geradas)} imagem(ns) salva(s) em imagens_cache/{caminho.stem}/")

        paginas_com_texto = [
            f"[Página {n}]\n{t}"
            for n, t in paginas
            if len(t) >= MIN_CHARS_PAGINA
        ]

        if not paginas_com_texto:
            print("  [AVISO] Nenhuma página com texto suficiente. Pulando chunking.")
            return 0

        texto_total = "\n\n".join(paginas_com_texto)

    # Etapa 3 — Chunking
    chunks = chunk_texto(texto_total)
    print(f"  {len(chunks)} chunk(s) gerado(s).")

    # Etapa 4 — Embeddings e inserção no banco
    print(f"  Gerando embeddings e salvando no Supabase...")
    total_salvos = 0

    for i, chunk in enumerate(chunks):
        embedding = gerar_embedding_documento(chunk)

        row: dict = {
            "tipo":         tipo,
            "nome_arquivo": caminho.name,
            "chunk_index":  i,
            "conteudo":     chunk,
            "embedding":    embedding,
        }
        if grupo:
            row["grupo"] = grupo
        supabase.table("documentos").insert(row).execute()

        total_salvos += 1
        print(f"  Chunk {total_salvos}/{len(chunks)} salvo.", end="\r")

    print(f"\n  Concluído: {total_salvos} chunk(s) indexado(s).          ")
    return total_salvos


def indexar_bytes(
    nome_arquivo: str,
    conteudo:     bytes,
    tipo:         str,
    relatorio:    str | None = None,
    grupo:        str | None = None,
) -> int:
    """
    Versão do indexar_documento que aceita bytes diretamente (para uso via API).

    Salva o PDF temporariamente, executa o pipeline de ingestão e remove o arquivo.
    Se relatorio (JSON string) for fornecido, indexa também como chunk especial.

    Args:
        nome_arquivo: Nome original do arquivo.
        conteudo:     Bytes do arquivo.
        tipo:         'manual' ou 'referencia'.
        relatorio:    JSON do relatório de revisão (opcional, apenas para referências aprovadas).
        grupo:        Nome do grupo organizacional (opcional, ex: 'Churrasqueiras').

    Returns:
        Número total de chunks salvos.
    """
    import tempfile
    import os

    # Usa apenas o nome do arquivo (sem subpastas) — uploads de pasta enviam caminhos relativos
    nome_base = Path(nome_arquivo).name
    ext = Path(nome_base).suffix.lower() or ".pdf"
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False, prefix="rag_") as tmp:
        tmp.write(conteudo)
        tmp_path = Path(tmp.name)

    # Renomeia para preservar o nome do arquivo (necessário para o cache de imagens)
    destino = tmp_path.parent / nome_base
    tmp_path.rename(destino)

    try:
        total = indexar_documento(destino, tipo, grupo)

        # Se veio com relatório de aprovação, indexa o conteúdo do relatório também
        if relatorio and tipo == "referencia":
            supabase = create_client(SUPABASE_URL, SUPABASE_KEY)
            embedding = gerar_embedding_documento(relatorio[:2000])

            supabase.table("documentos").insert({
                "tipo":         "referencia",
                "nome_arquivo": nome_arquivo,
                "chunk_index":  9999,  # índice especial para identificar o relatório
                "conteudo":     f"[RELATÓRIO DE REVISÃO APROVADO]\n{relatorio[:2000]}",
                "embedding":    embedding,
            }).execute()

            total += 1
            print("  Relatório de aprovação indexado.")

        return total

    finally:
        # Garante remoção do arquivo temporário
        if destino.exists():
            os.unlink(destino)


# ---------------------------------------------------------------------------
# Interface de linha de comando
# ---------------------------------------------------------------------------

def _indexar_diretorio(tipo: str) -> None:
    """Indexa todos os PDFs e DOCXs do diretório correspondente ao tipo (recursivo)."""
    pasta = DIR_MANUAL if tipo == "manual" else DIR_PROJETOS
    arquivos = sorted(
        list(pasta.rglob("*.pdf")) + list(pasta.rglob("*.docx"))
    )

    if not arquivos:
        print(f"Nenhum PDF ou DOCX encontrado em: {pasta}")
        print(f"Adicione arquivos nessa pasta e execute novamente.")
        return

    print(f"Encontrados {len(arquivos)} arquivo(s) em {pasta.name}/ (incluindo subpastas)")
    total_geral = 0

    for caminho in arquivos:
        total_geral += indexar_documento(caminho, tipo)

    print(f"\n{'='*60}")
    print(f"Ingestão concluída. Total: {total_geral} chunk(s) no banco.")
    print(f"{'='*60}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Indexa documentos da base de contexto no Supabase.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Exemplos:\n"
            "  python backend/ingestao.py --tipo manual\n"
            "  python backend/ingestao.py --tipo referencia\n"
        ),
    )
    parser.add_argument(
        "--tipo",
        choices=["manual", "referencia"],
        required=True,
        help="Tipo de documento a indexar.",
    )
    args = parser.parse_args()

    try:
        validar_config()
        configurar_gemini()
        _indexar_diretorio(args.tipo)
    except EnvironmentError as e:
        print(f"\nErro de configuração: {e}")
        sys.exit(1)
    except KeyboardInterrupt:
        print("\nIngestão interrompida pelo usuário.")
        sys.exit(0)
